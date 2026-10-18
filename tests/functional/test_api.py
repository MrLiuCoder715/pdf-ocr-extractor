from io import BytesIO

import pytest
from openpyxl import load_workbook

from pdf_ocr_extractor.config import Settings
from pdf_ocr_extractor.excel_writer import SHEET_TITLE
from pdf_ocr_extractor.models import ExtractionResult
from pdf_ocr_extractor.pdf_renderer import InvalidPdfError
from pdf_ocr_extractor.web import create_app

PDF = b"%PDF-1.7 fake body"


class FakeService:
    def __init__(self):
        self.calls = []
        self.column_calls = []

    def extract(self, pdf_bytes, filename, titles, column_titles=()):
        self.calls.append((filename, list(titles)))
        self.column_calls.append(list(column_titles))
        if b"broken" in pdf_bytes:
            raise InvalidPdfError("bad")
        fields = {t: f"{t}-value" for t in titles if t not in column_titles}
        rows = [{t: f"{t}-{i}" for t in column_titles} for i in (1, 2)] if column_titles else []
        return ExtractionResult(filename, fields, rows)


@pytest.fixture
def service():
    return FakeService()


@pytest.fixture
def client(service):
    app = create_app(service=service, settings=Settings(max_upload_mb=1))
    app.testing = True
    return app.test_client()


def post_extract(client, files, titles, column_titles=()):
    data = {
        "titles": titles,
        "column_titles": list(column_titles),
        "files": [(BytesIO(body), name) for name, body in files],
    }
    return client.post("/api/extract", data=data, content_type="multipart/form-data")


def test_index_serves_ui(client):
    resp = client.get("/")
    assert resp.status_code == 200
    assert "PDF" in resp.get_data(as_text=True)


def test_extract_success(client, service):
    resp = post_extract(client, [("a.pdf", PDF), ("b.PDF", PDF)], ["工程編號", "日期"])
    assert resp.status_code == 200
    body = resp.get_json()
    assert body["titles"] == ["工程編號", "日期"]
    assert [r["filename"] for r in body["results"]] == ["a.pdf", "b.PDF"]
    assert body["results"][0]["fields"] == {"工程編號": "工程編號-value", "日期": "日期-value"}


def test_extract_cleans_titles(client, service):
    resp = post_extract(client, [("a.pdf", PDF)], [" 工程編號 ", "", "工程編號", "日期"])
    assert resp.get_json()["titles"] == ["工程編號", "日期"]
    assert service.calls == [("a.pdf", ["工程編號", "日期"])]


def test_extract_keeps_only_basename(client):
    resp = post_extract(client, [("scan/工程.pdf", PDF)], ["工程編號"])
    assert resp.get_json()["results"][0]["filename"] == "工程.pdf"


def test_extract_requires_files(client):
    resp = client.post("/api/extract", data={"titles": ["工程編號"]}, content_type="multipart/form-data")
    assert resp.status_code == 400
    assert "error" in resp.get_json()


def test_extract_requires_titles(client):
    resp = post_extract(client, [("a.pdf", PDF)], ["  "])
    assert resp.status_code == 400
    assert "error" in resp.get_json()


@pytest.mark.parametrize("name,body", [("note.txt", PDF), ("fake.pdf", b"hello world")])
def test_extract_rejects_non_pdf(client, service, name, body):
    resp = post_extract(client, [(name, body)], ["工程編號"])
    assert resp.status_code == 400
    assert name in resp.get_json()["error"]
    assert service.calls == []


def test_extract_reports_unreadable_pdf(client):
    resp = post_extract(client, [("x.pdf", PDF + b"broken")], ["工程編號"])
    assert resp.status_code == 400
    assert "x.pdf" in resp.get_json()["error"]


def test_extract_rejects_oversized_upload(client):
    big = PDF + b"0" * (1024 * 1024 + 10)
    resp = post_extract(client, [("big.pdf", big)], ["工程編號"])
    assert resp.status_code == 413
    assert "error" in resp.get_json()


def test_export_returns_xlsx(client):
    payload = {
        "titles": ["工程編號"],
        "results": [{"filename": "a.pdf", "fields": {"工程編號": "nglc5zxxxd"}}],
    }
    resp = client.post("/api/export", json=payload)
    assert resp.status_code == 200
    assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in resp.headers["Content-Disposition"]
    ws = load_workbook(BytesIO(resp.data))[SHEET_TITLE]
    assert [c.value for c in ws[2]] == ["a.pdf", "nglc5zxxxd"]


@pytest.mark.parametrize(
    "payload",
    [
        None,
        {"titles": "工程編號", "results": []},
        {"titles": ["工程編號"], "results": [{"filename": 1, "fields": {}}]},
        {"titles": ["工程編號"], "results": [{"filename": "a.pdf", "fields": []}]},
    ],
)
def test_export_rejects_bad_payload(client, payload):
    resp = client.post("/api/export", data="oops", content_type="application/json") if payload is None \
        else client.post("/api/export", json=payload)
    assert resp.status_code == 400
    assert "error" in resp.get_json()


def test_extract_returns_rows_and_expanded_table(client, service):
    resp = post_extract(client, [("a.pdf", PDF)], ["承商名稱", "姓名"], ["姓名", "不存在"])
    assert resp.status_code == 200
    body = resp.get_json()
    assert body["column_titles"] == ["姓名"]
    assert service.column_calls == [["姓名"]]
    assert body["results"][0]["rows"] == [{"姓名": "姓名-1"}, {"姓名": "姓名-2"}]
    assert body["table"] == {
        "header": ["檔名", "承商名稱", "姓名"],
        "rows": [["a.pdf", "承商名稱-value", "姓名-1"], ["a.pdf", "承商名稱-value", "姓名-2"]],
    }


def test_extract_without_column_titles_has_one_row_per_pdf(client):
    body = post_extract(client, [("a.pdf", PDF), ("b.pdf", PDF)], ["工程編號"]).get_json()
    assert body["column_titles"] == []
    assert body["table"]["rows"] == [["a.pdf", "工程編號-value"], ["b.pdf", "工程編號-value"]]


def test_export_expands_table_rows(client):
    payload = {
        "titles": ["承商名稱", "姓名"],
        "column_titles": ["姓名"],
        "results": [
            {"filename": "a.pdf", "fields": {"承商名稱": "樺佑"}, "rows": [{"姓名": "王小明"}, {"姓名": "陳大華"}]}
        ],
    }
    resp = client.post("/api/export", json=payload)
    assert resp.status_code == 200
    ws = load_workbook(BytesIO(resp.data))[SHEET_TITLE]
    assert [[c.value for c in r] for r in ws.iter_rows(min_row=2)] == [
        ["a.pdf", "樺佑", "王小明"],
        ["a.pdf", "樺佑", "陳大華"],
    ]


@pytest.mark.parametrize(
    "payload",
    [
        {"titles": ["姓名"], "column_titles": "姓名", "results": []},
        {"titles": ["姓名"], "results": [{"filename": "a.pdf", "fields": {}, "rows": {}}]},
        {"titles": ["姓名"], "results": [{"filename": "a.pdf", "fields": {}, "rows": ["x"]}]},
    ],
)
def test_export_rejects_bad_table_payload(client, payload):
    resp = client.post("/api/export", json=payload)
    assert resp.status_code == 400
