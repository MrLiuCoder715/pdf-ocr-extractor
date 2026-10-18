"""Browser E2E tests. Requires the app running (docker compose up) — see README.

Run: uv run pytest -m e2e
"""

import os
from io import BytesIO

import pytest
from openpyxl import load_workbook
from playwright.sync_api import Page, expect

from tests.conftest import ROSTER_NAMES, make_pdf, make_roster_pdf

pytestmark = pytest.mark.e2e

BASE_URL = os.environ.get("E2E_BASE_URL", "http://localhost:8000")
OCR_TIMEOUT_MS = 120_000


@pytest.fixture
def ui(page: Page) -> Page:
    page.goto(BASE_URL)
    page.evaluate("localStorage.clear()")
    page.reload()
    return page


def title_items(page: Page):
    return page.get_by_role("list", name="標題清單").get_by_role("listitem")


def add_title(page: Page, title: str) -> None:
    page.get_by_label("新標題").fill(title)
    page.get_by_role("button", name="新增").click()


def test_default_title_present(ui: Page):
    expect(title_items(ui)).to_have_count(1)
    expect(title_items(ui).first).to_contain_text("工程編號")


def test_add_delete_and_persist_titles(ui: Page):
    add_title(ui, "日期")
    expect(title_items(ui)).to_have_count(2)

    add_title(ui, "工程編號")
    expect(ui.get_by_role("status")).to_contain_text("已存在")
    expect(title_items(ui)).to_have_count(2)

    ui.reload()
    expect(title_items(ui)).to_have_count(2)
    expect(title_items(ui).nth(1)).to_contain_text("日期")

    ui.get_by_role("button", name="刪除 日期").click()
    expect(title_items(ui)).to_have_count(1)


def test_extract_without_file_shows_error(ui: Page):
    ui.get_by_role("button", name="開始擷取").click()
    expect(ui.get_by_role("status")).to_contain_text("請先選擇 PDF")


def test_upload_extract_and_download_excel(ui: Page, tmp_path):
    pdf_path = tmp_path / "工程文件.pdf"
    pdf_path.write_bytes(make_pdf([["工程編號：nglc5zxxxd", "日期：2026/10/16"]]))
    add_title(ui, "日期")

    ui.get_by_label("選擇 PDF").set_input_files(pdf_path)
    ui.get_by_role("button", name="開始擷取").click()

    table = ui.locator("#result-table")
    expect(table).to_contain_text("nglc5zxxxd", timeout=OCR_TIMEOUT_MS)
    expect(table).to_contain_text("工程文件.pdf")
    expect(table).to_contain_text("2026/10/16")

    with ui.expect_download() as info:
        ui.get_by_role("button", name="下載 Excel").click()
    target = tmp_path / info.value.suggested_filename
    info.value.save_as(target)

    ws = load_workbook(BytesIO(target.read_bytes())).active
    assert [c.value for c in ws[1]] == ["檔名", "工程編號", "日期"]
    assert [c.value for c in ws[2]] == ["工程文件.pdf", "nglc5zxxxd", "2026/10/16"]


def test_old_saved_titles_are_migrated_to_field_mode(page: Page):
    page.goto(BASE_URL)
    page.evaluate("localStorage.setItem('pdf-ocr-extractor.titles', JSON.stringify(['工程編號', '日期']))")
    page.reload()
    expect(title_items(page)).to_have_count(2)
    expect(page.get_by_role("button", name="切換 日期 模式")).to_have_text("欄位")


def test_toggle_mode_persists(ui: Page):
    add_title(ui, "姓名")
    toggle = ui.get_by_role("button", name="切換 姓名 模式")
    expect(toggle).to_have_text("欄位")
    toggle.click()
    expect(toggle).to_have_text("表格欄")
    ui.reload()
    expect(ui.get_by_role("button", name="切換 姓名 模式")).to_have_attribute("aria-pressed", "true")


def test_roster_table_column_one_row_per_name(ui: Page, tmp_path):
    pdf_path = tmp_path / "施工人員名冊.pdf"
    pdf_path.write_bytes(make_roster_pdf())
    ui.get_by_role("button", name="刪除 工程編號").click()
    add_title(ui, "承商名稱")
    add_title(ui, "姓名")
    ui.get_by_role("button", name="切換 姓名 模式").click()

    ui.get_by_label("選擇 PDF").set_input_files(pdf_path)
    ui.get_by_role("button", name="開始擷取").click()

    table = ui.locator("#result-table")
    expect(table.locator("tr")).to_have_count(1 + len(ROSTER_NAMES), timeout=OCR_TIMEOUT_MS)
    for name in ROSTER_NAMES:
        expect(table).to_contain_text(name)

    with ui.expect_download() as info:
        ui.get_by_role("button", name="下載 Excel").click()
    target = tmp_path / info.value.suggested_filename
    info.value.save_as(target)

    ws = load_workbook(BytesIO(target.read_bytes())).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert rows[0] == ["檔名", "承商名稱", "姓名"]
    assert rows[1:] == [["施工人員名冊.pdf", "測試公司", name] for name in ROSTER_NAMES]
