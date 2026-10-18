from io import BytesIO

from openpyxl import load_workbook

from pdf_ocr_extractor.excel_writer import SHEET_TITLE, write_workbook

HEADER = ["檔名", "工程編號", "日期"]
ROWS = [["a.pdf", "A1", "2026/10/16"], ["b.pdf", "B2", ""]]


def load(data: bytes):
    return load_workbook(BytesIO(data))[SHEET_TITLE]


def test_header_and_rows():
    ws = load(write_workbook(HEADER, ROWS))
    rows = [[c if c is not None else "" for c in r] for r in ws.iter_rows(values_only=True)]
    assert rows == [HEADER, *ROWS]


def test_header_bold_and_frozen():
    ws = load(write_workbook(HEADER, ROWS))
    assert all(c.font.bold for c in ws[1])
    assert ws.freeze_panes == "A2"


def test_values_kept_as_text():
    ws = load(write_workbook(["檔名", "編號"], [["c.pdf", "00123"]]))
    assert ws["B2"].value == "00123"
    assert ws["B2"].data_type == "s"


def test_blank_value_is_empty_cell():
    ws = load(write_workbook(["檔名", "編號"], [["d.pdf", ""]]))
    assert ws["B2"].value in (None, "")


def test_formula_like_values_stored_as_text():
    ws = load(write_workbook(["檔名", "編號"], [["=evil.pdf", "=1+1"]]))
    assert ws["A2"].data_type == "s" and ws["A2"].value == "=evil.pdf"
    assert ws["B2"].data_type == "s" and ws["B2"].value == "=1+1"
