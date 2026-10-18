"""Write a header + rows table to an .xlsx workbook."""

from collections.abc import Sequence
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

SHEET_TITLE = "擷取結果"


def _append_text_row(ws, values: Sequence[str]) -> None:
    ws.append(list(values))
    for cell in ws[ws.max_row]:
        if cell.value not in (None, ""):
            # Force string type so OCR text like "=..." is never evaluated as a formula.
            cell.data_type = "s"


def write_workbook(header: Sequence[str], rows: Sequence[Sequence[str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE
    _append_text_row(ws, header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for row in rows:
        _append_text_row(ws, row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
